from pathlib import Path
import sys
import openpyxl

# Support running both as `python -m backend.seed_data` and `python backend/seed_data.py`
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from backend.database import get_db_connection, init_db

SEED_FILE_PATH = BASE_DIR / "data" / "events_seed.xlsx"


def seed_events() -> int:
    """
    Seed the events table from data/events_seed.xlsx ("Events" sheet).
    Skips seeding if the events table already contains rows.
    Raises FileNotFoundError with a clear message if data/events_seed.xlsx is missing.
    """
    init_db()
    conn = get_db_connection()
    try:
        cursor = conn.cursor()
        cursor.execute("SELECT COUNT(*) FROM events;")
        existing_count = cursor.fetchone()[0]
        if existing_count > 0:
            return 0

        if not SEED_FILE_PATH.exists():
            raise FileNotFoundError(
                f"Missing seed file at '{SEED_FILE_PATH}'. "
                "Please copy the provided 'events_seed.xlsx' file into the 'data/' folder "
                "before running this script or starting the server."
            )

        wb = openpyxl.load_workbook(SEED_FILE_PATH, data_only=True)
        if "Events" in wb.sheetnames:
            sheet = wb["Events"]
        else:
            sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) <= 1:
            raise ValueError("The 'Events' sheet in data/events_seed.xlsx contains no event rows.")

        # Row 1 is the header row:
        # id, title, category, description, event_date, event_time, venue, capacity, fee, icon
        inserted = 0
        for row in rows[1:]:
            if not row or all(cell is None for cell in row):
                continue

            event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = (
                row[0],
                str(row[1]).strip() if row[1] is not None else "",
                str(row[2]).strip() if row[2] is not None else "",
                str(row[3]).strip() if row[3] is not None else "",
                str(row[4]).strip() if row[4] is not None else "",
                str(row[5]).strip() if row[5] is not None else "",
                str(row[6]).strip() if row[6] is not None else "",
                int(row[7]) if row[7] is not None else 0,
                int(row[8]) if row[8] is not None else 0,
                str(row[9]).strip() if len(row) > 9 and row[9] is not None else "✨",
            )

            # Format date cleanly if openpyxl parsed it as a datetime/date object
            if hasattr(row[4], "strftime"):
                event_date = row[4].strftime("%Y-%m-%d")
            if hasattr(row[5], "strftime"):
                event_time = row[5].strftime("%H:%M")

            if event_id is not None:
                cursor.execute(
                    """
                    INSERT INTO events (
                        id, title, category, description, event_date,
                        event_time, venue, capacity, fee, icon
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """,
                    (
                        int(event_id),
                        title,
                        category,
                        description,
                        event_date,
                        event_time,
                        venue,
                        capacity,
                        fee,
                        icon,
                    ),
                )
            else:
                cursor.execute(
                    """
                    INSERT INTO events (
                        title, category, description, event_date,
                        event_time, venue, capacity, fee, icon
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """,
                    (
                        title,
                        category,
                        description,
                        event_date,
                        event_time,
                        venue,
                        capacity,
                        fee,
                        icon,
                    ),
                )
            inserted += 1

        conn.commit()
        return inserted
    finally:
        conn.close()


def print_seeded_events() -> None:
    """Print all rows in the events table for verification."""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    conn = get_db_connection()
    try:
        rows = conn.execute("SELECT * FROM events ORDER BY id ASC;").fetchall()
        print(f"\nTotal events in database: {len(rows)}")
        print("-" * 100)
        for r in rows:
            print(
                f"[{r['id']:02d}] {r['icon']} {r['title']} ({r['category']}) | "
                f"{r['event_date']} {r['event_time']} @ {r['venue']} | "
                f"Cap: {r['capacity']} | Fee: Rs.{r['fee']}"
            )
        print("-" * 100)
    finally:
        conn.close()


if __name__ == "__main__":
    count = seed_events()
    if count > 0:
        print(f"Successfully seeded {count} events from {SEED_FILE_PATH.name}.")
    else:
        print("Events table already contains rows; skipping seed.")
    print_seeded_events()
